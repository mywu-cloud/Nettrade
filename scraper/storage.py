"""寫檔：前端用 JSON、封存用 CSV + Excel、日期索引。"""
from __future__ import annotations

import csv
import json
import logging
from pathlib import Path

from .models import Board
from .utils import now_tw

log = logging.getLogger("scraper.storage")

RED, GREEN = "D92B2B", "1A9E3F"


def write_day(boards: list[Board], day: str, cfg: dict, root: Path, demo: bool = False) -> Path:
    out_cfg = cfg["output"]
    web_dir = root / out_cfg["web_data_dir"]
    year_dir = web_dir / day[:4]
    year_dir.mkdir(parents=True, exist_ok=True)

    payload = {
        "date": day,
        "generated_at": now_tw().isoformat(timespec="seconds"),
        "demo": demo,
        "unit": "張",
        "boards": [b.to_dict() for b in boards],
    }
    day_file = year_dir / f"{day}.json"
    _write_json(day_file, payload)
    _write_json(web_dir / "latest.json", payload)
    _update_index(web_dir, day)

    archive = root / out_cfg["archive_dir"] / day[:4]
    if out_cfg.get("write_csv", True):
        _write_csv(boards, archive / f"{day}.csv")
    if out_cfg.get("write_xlsx", True):
        _write_xlsx(boards, archive / f"{day}.xlsx", day)
    log.info("已寫入 %s（%d 張排行表）", day_file, len(boards))
    return day_file


def _write_json(path: Path, obj) -> None:
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(obj, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    tmp.replace(path)                      # 原子寫入，前端不會讀到半個檔


def _update_index(web_dir: Path, day: str) -> None:
    idx_path = web_dir / "index.json"
    dates: list[str] = []
    if idx_path.exists():
        dates = json.loads(idx_path.read_text(encoding="utf-8")).get("dates", [])
    dates = sorted(set(dates) | {day}, reverse=True)
    _write_json(idx_path, {"dates": dates, "updated": now_tw().isoformat(timespec="seconds")})


def _write_csv(boards: list[Board], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:   # BOM 讓 Excel 正確顯示中文
        w = csv.writer(f)
        w.writerow(["來源", "類別", "市場", "方向", "名次", "代號", "名稱", "買賣超(張)", "收盤價", "漲跌", "漲跌幅%"])
        for b in boards:
            for side, rows in (("買超", b.buy), ("賣超", b.sell)):
                for r in rows:
                    w.writerow([b.source_name, b.category_name, b.market_name, side, r.rank, r.code, r.name,
                                r.net, r.close, r.change, r.change_pct])


def _write_xlsx(boards: list[Board], path: Path, day: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="2F3B52")
    head_font = Font(bold=True, color="FFFFFF")

    for b in boards:
        if not b.buy and not b.sell:
            continue
        ws = wb.create_sheet(f"{b.source_name[:4]}_{b.category_name}_{b.market_name}"[:31])
        ws["A1"] = f"{b.source_name}｜{b.market_name}{b.category_name}買賣超排行｜{b.data_date or day}｜單位：張"
        ws["A1"].font = Font(bold=True, size=13)
        headers = ["名次", "代號", "名稱", "買超(張)", "收盤價", "漲跌"]
        for side_i, (title, rows, color) in enumerate((("買超", b.buy, RED), ("賣超", b.sell, GREEN))):
            c0 = 1 + side_i * 7
            ws.cell(row=2, column=c0, value=title).font = Font(bold=True, color=color, size=12)
            for j, h in enumerate(headers):
                cell = ws.cell(row=3, column=c0 + j, value=h if j != 3 else f"{title}(張)")
                cell.fill, cell.font = head_fill, head_font
                cell.alignment = Alignment(horizontal="center")
            for i, r in enumerate(rows):
                vals = [r.rank, r.code, r.name, r.net, r.close, r.change]
                for j, v in enumerate(vals):
                    cell = ws.cell(row=4 + i, column=c0 + j, value=v)
                    if j == 3:
                        cell.number_format = "#,##0"
                        cell.font = Font(bold=True, color=RED if r.net > 0 else GREEN)
                    if j == 5 and isinstance(v, (int, float)) and v:
                        cell.font = Font(color=RED if v > 0 else GREEN)
        for col in range(1, 14):
            ws.column_dimensions[get_column_letter(col)].width = 12 if col % 7 != 3 else 16
        ws.freeze_panes = "A4"

    ws = wb.create_sheet("說明")
    ws["A1"] = "本檔由程式自動抓取公開網頁整理，僅供參考，不構成投資建議。"
    ws["A2"] = "顏色慣例：紅色＝買超／上漲，綠色＝賣超／下跌。"
    wb.save(path)
